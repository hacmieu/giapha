"""
Views cho hệ thống Tự Khai Gia Phả (Draft Submission).
- Public views: không cần đăng nhập, dùng session + IP xác minh
- Admin views: cần login + quyền editor
"""
from django.views.generic import TemplateView, CreateView, UpdateView, DeleteView, DetailView, ListView
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.shortcuts import get_object_or_404, redirect, render
from django.contrib import messages
from django.http import JsonResponse, HttpResponseForbidden
from django.db import models
from django.utils import timezone
from django.db.models import Count
import json

from .models import DraftSubmission, DraftPerson, DraftSpouseRelation, FamilyMember, Chi, SpouseRelation
from .draft_forms import DraftSubmissionForm, DraftPersonForm, DraftLookupForm


# ═══════════════════════════════════════════════════════════════
# HELPERS
# ═══════════════════════════════════════════════════════════════

def _get_client_ip(request):
    x_forwarded = request.META.get('HTTP_X_FORWARDED_FOR')
    if x_forwarded:
        return x_forwarded.split(',')[0].strip()
    return request.META.get('REMOTE_ADDR', '0.0.0.0')


def _can_access_draft(request, draft):
    """Kiểm tra người dùng có quyền truy cập bản khai không"""
    # Admin luôn được truy cập
    if request.user.is_authenticated and request.user.can_edit_members():
        return True
    # Mã draft_code đã được lưu trong session (từ lần lookup hoặc tạo)
    my_codes = request.session.get('my_draft_codes', [])
    if draft.draft_code in my_codes:
        return True
    # Session gốc khớp
    if request.session.session_key and request.session.session_key == draft.session_key:
        return True
    # IP khớp
    if _get_client_ip(request) == draft.ip_address:
        return True
    return False


def _can_edit_draft(request, draft):
    """Chỉ chủ bản khai mới sửa, và chỉ khi status=draft"""
    if draft.status != 'draft':
        return False
    return _can_access_draft(request, draft)


def _remember_draft(request, draft_code):
    """Lưu mã bản khai vào session để lần sau mở lại được"""
    if not request.session.session_key:
        request.session.create()
    codes = request.session.get('my_draft_codes', [])
    if draft_code not in codes:
        codes.append(draft_code)
        request.session['my_draft_codes'] = codes
        request.session.modified = True


MAX_DRAFTS_PER_IP_PER_DAY = 5
MAX_PERSONS_PER_DRAFT = 50


# ═══════════════════════════════════════════════════════════════
# PUBLIC VIEWS — Không cần đăng nhập
# ═══════════════════════════════════════════════════════════════

class DraftLandingView(TemplateView):
    """Trang giới thiệu + tạo bản khai mới / tra cứu"""
    template_name = 'drafts/landing.html'

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['lookup_form'] = DraftLookupForm()
        ctx['create_form'] = DraftSubmissionForm()

        # Hiện bản khai gần đây từ session
        my_codes = self.request.session.get('my_draft_codes', [])
        if my_codes:
            ctx['my_drafts'] = DraftSubmission.objects.filter(
                draft_code__in=my_codes
            ).order_by('-updated_at')
        return ctx


class DraftCreateView(CreateView):
    """Tạo bản khai mới"""
    model = DraftSubmission
    form_class = DraftSubmissionForm
    template_name = 'drafts/landing.html'

    def form_valid(self, form):
        ip = _get_client_ip(self.request)

        # Rate limit
        today_count = DraftSubmission.objects.filter(
            ip_address=ip,
            created_at__date=timezone.now().date()
        ).count()
        if today_count >= MAX_DRAFTS_PER_IP_PER_DAY:
            messages.error(self.request, f'Bạn đã tạo quá {MAX_DRAFTS_PER_IP_PER_DAY} bản khai hôm nay. Vui lòng quay lại ngày mai.')
            return redirect('draft_landing')

        # Ensure session exists
        if not self.request.session.session_key:
            self.request.session.create()

        draft = form.save(commit=False)
        draft.ip_address = ip
        draft.session_key = self.request.session.session_key
        draft.save()

        # Nếu người khai nằm trong gia phả → tự thêm làm người #1
        is_in_family = form.cleaned_data.get('is_in_family', 'yes')
        if is_in_family == 'yes':
            gender = form.cleaned_data.get('submitter_gender', 'male') or 'male'
            DraftPerson.objects.create(
                submission=draft,
                temp_id=1,
                name=draft.submitter_name,
                gender=gender,
                relation_to_submitter='self',
            )
            msg = f'Đã tạo bản khai {draft.draft_code}! Bạn ({draft.submitter_name}) là người #1. Hãy thêm người thân tiếp.'
        else:
            msg = f'Đã tạo bản khai {draft.draft_code}! Hãy thêm người thân vào danh sách.'

        _remember_draft(self.request, draft.draft_code)
        messages.success(self.request, msg)
        return redirect('draft_detail', draft_code=draft.draft_code)

    def form_invalid(self, form):
        messages.error(self.request, 'Vui lòng điền đủ thông tin.')
        return redirect('draft_landing')


class DraftLookupView(TemplateView):
    """Tra cứu bản khai bằng mã — mã đúng = được truy cập"""
    template_name = 'drafts/landing.html'

    def post(self, request, *args, **kwargs):
        code = request.POST.get('draft_code', '').strip().upper()
        # Normalize: accept with or without DK- prefix
        if not code.startswith('DK-'):
            code = f'DK-{code}'
        try:
            draft = DraftSubmission.objects.get(draft_code__iexact=code)
            # Có mã = có quyền → lưu vào session
            _remember_draft(request, draft.draft_code)
            return redirect('draft_detail', draft_code=draft.draft_code)
        except DraftSubmission.DoesNotExist:
            messages.error(request, f'Không tìm thấy bản khai mã "{code}".')
        return redirect('draft_landing')


class DraftDetailView(DetailView):
    """Xem chi tiết bản khai + danh sách người"""
    model = DraftSubmission
    template_name = 'drafts/detail.html'
    context_object_name = 'draft'
    slug_field = 'draft_code'
    slug_url_kwarg = 'draft_code'

    def dispatch(self, request, *args, **kwargs):
        draft = self.get_object()
        if not _can_access_draft(request, draft):
            return HttpResponseForbidden('Bạn không có quyền xem bản khai này.')
        # Ghi nhớ draft_code vào session để lần sau mở lại được
        _remember_draft(request, draft.draft_code)
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['persons'] = self.object.persons.order_by('temp_id')
        ctx['can_edit'] = _can_edit_draft(self.request, self.object)
        ctx['person_form'] = DraftPersonForm(submission=self.object)
        return ctx


class DraftAddPersonView(CreateView):
    """Thêm người vào bản khai"""
    model = DraftPerson
    form_class = DraftPersonForm
    template_name = 'drafts/person_form.html'

    def dispatch(self, request, *args, **kwargs):
        self.draft = get_object_or_404(DraftSubmission, draft_code=kwargs['draft_code'])
        if not _can_edit_draft(request, self.draft):
            messages.error(request, 'Không thể sửa bản khai này.')
            return redirect('draft_detail', draft_code=self.draft.draft_code)
        if self.draft.persons.count() >= MAX_PERSONS_PER_DRAFT:
            messages.error(request, f'Tối đa {MAX_PERSONS_PER_DRAFT} người mỗi bản khai.')
            return redirect('draft_detail', draft_code=self.draft.draft_code)
        return super().dispatch(request, *args, **kwargs)

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['submission'] = self.draft
        return kwargs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['draft'] = self.draft
        ctx['title'] = 'Thêm người thân'
        ctx['is_create'] = True
        return ctx

    def form_valid(self, form):
        person = form.save(commit=False)
        person.submission = self.draft
        # Auto-assign temp_id
        last = self.draft.persons.order_by('-temp_id').first()
        person.temp_id = (last.temp_id + 1) if last else 1
        person.save()
        messages.success(self.request, f'Đã thêm: {person.name}')
        return redirect('draft_detail', draft_code=self.draft.draft_code)

    def form_invalid(self, form):
        return self.render_to_response(self.get_context_data(form=form))


class DraftEditPersonView(UpdateView):
    """Sửa thông tin người trong bản khai"""
    model = DraftPerson
    form_class = DraftPersonForm
    template_name = 'drafts/person_form.html'
    pk_url_kwarg = 'person_pk'

    def dispatch(self, request, *args, **kwargs):
        self.draft = get_object_or_404(DraftSubmission, draft_code=kwargs['draft_code'])
        if not _can_edit_draft(request, self.draft):
            messages.error(request, 'Không thể sửa bản khai này.')
            return redirect('draft_detail', draft_code=self.draft.draft_code)
        return super().dispatch(request, *args, **kwargs)

    def get_queryset(self):
        return DraftPerson.objects.filter(submission=self.draft)

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['submission'] = self.draft
        return kwargs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['draft'] = self.draft
        ctx['title'] = f'Sửa: {self.object.name}'
        ctx['is_create'] = False
        return ctx

    def form_valid(self, form):
        form.save()
        messages.success(self.request, f'Đã cập nhật: {self.object.name}')
        return redirect('draft_detail', draft_code=self.draft.draft_code)


class DraftDeletePersonView(DeleteView):
    """Xóa người trong bản khai"""
    model = DraftPerson
    pk_url_kwarg = 'person_pk'

    def dispatch(self, request, *args, **kwargs):
        self.draft = get_object_or_404(DraftSubmission, draft_code=kwargs['draft_code'])
        if not _can_edit_draft(request, self.draft):
            messages.error(request, 'Không thể sửa bản khai này.')
            return redirect('draft_detail', draft_code=self.draft.draft_code)
        return super().dispatch(request, *args, **kwargs)

    def get_queryset(self):
        return DraftPerson.objects.filter(submission=self.draft)

    def post(self, request, *args, **kwargs):
        person = self.get_object()
        name = person.name
        person.delete()
        messages.success(request, f'Đã xóa: {name}')
        return redirect('draft_detail', draft_code=self.draft.draft_code)


class DraftSubmitView(TemplateView):
    """Gửi bản khai (lock lại)"""
    template_name = 'drafts/detail.html'

    def post(self, request, draft_code):
        draft = get_object_or_404(DraftSubmission, draft_code=draft_code)
        if not _can_edit_draft(request, draft):
            messages.error(request, 'Không thể gửi bản khai này.')
            return redirect('draft_detail', draft_code=draft_code)

        if draft.persons.count() == 0:
            messages.error(request, 'Bản khai chưa có người nào. Hãy thêm ít nhất 1 người.')
            return redirect('draft_detail', draft_code=draft_code)

        draft.status = 'submitted'
        draft.save(update_fields=['status', 'updated_at'])
        messages.success(request, f'Đã gửi bản khai {draft.draft_code} cho trưởng họ duyệt! Lưu mã này để tra cứu.')
        return redirect('draft_detail', draft_code=draft_code)


class DraftRecallView(TemplateView):
    """Thu hồi bản khai đã gửi → quay lại trạng thái draft để sửa"""
    template_name = 'drafts/detail.html'

    def post(self, request, draft_code):
        draft = get_object_or_404(DraftSubmission, draft_code=draft_code)

        # Chỉ thu hồi được khi đang submitted (chưa duyệt/từ chối)
        if draft.status != 'submitted':
            messages.error(request, 'Chỉ có thể thu hồi bản khai đang chờ duyệt.')
            return redirect('draft_detail', draft_code=draft_code)

        # Kiểm tra quyền: chủ bản khai hoặc admin
        if not _can_access_draft(request, draft):
            messages.error(request, 'Bạn không có quyền thu hồi bản khai này.')
            return redirect('draft_detail', draft_code=draft_code)

        draft.status = 'draft'
        draft.save(update_fields=['status', 'updated_at'])
        messages.success(request, f'Đã thu hồi bản khai {draft.draft_code}. Bạn có thể sửa lại rồi gửi lại.')
        return redirect('draft_detail', draft_code=draft_code)


def draft_export_json(request, draft_code):
    """Xuất bản khai ra JSON"""
    draft = get_object_or_404(DraftSubmission, draft_code=draft_code)
    if not _can_access_draft(request, draft):
        return JsonResponse({'error': 'Không có quyền'}, status=403)

    persons = []
    for p in draft.persons.order_by('temp_id'):
        persons.append({
            'temp_id': p.temp_id,
            'name': p.name,
            'gender': p.gender,
            'relation_to_submitter': p.get_relation_to_submitter_display(),
            'birth_date': p.birth_date,
            'death_date': p.death_date,
            'birth_order': p.birth_order,
            'chi_number': p.chi_number,
            'generation': p.generation,
            'father_temp_id': p.father_temp_id,
            'mother_temp_id': p.mother_temp_id,
            'notes': p.notes,
        })

    # Spouse relations
    spouse_rels = []
    for rel in draft.spouse_relations.order_by('wife_order').all():
        spouse_rels.append({
            'person1_temp_id': rel.person1_temp_id,
            'person2_temp_id': rel.person2_temp_id,
            'status': rel.status,
            'wife_order': rel.wife_order,
        })

    data = {
        'draft_code': draft.draft_code,
        'submitter_name': draft.submitter_name,
        'submitter_phone': draft.submitter_phone,
        'submitter_relation': draft.submitter_relation,
        'status': draft.get_status_display(),
        'created_at': draft.created_at.isoformat(),
        'notes': draft.notes,
        'persons': persons,
        'spouse_relations': spouse_rels,
    }
    return JsonResponse(data, json_dumps_params={'ensure_ascii': False, 'indent': 2})


# ═══════════════════════════════════════════════════════════════
# EXCEL UPLOAD + VISUAL RELATIONSHIP BUILDER
# ═══════════════════════════════════════════════════════════════

import openpyxl
from io import BytesIO


def draft_upload_excel(request, draft_code):
    """Upload Excel để tạo danh sách người nhanh"""
    draft = get_object_or_404(DraftSubmission, draft_code=draft_code)
    if not _can_edit_draft(request, draft):
        messages.error(request, 'Không thể sửa bản khai này.')
        return redirect('draft_detail', draft_code=draft_code)

    if request.method != 'POST' or 'excel_file' not in request.FILES:
        messages.error(request, 'Vui lòng chọn file Excel.')
        return redirect('draft_detail', draft_code=draft_code)

    f = request.FILES['excel_file']
    if not f.name.endswith(('.xlsx', '.xls')):
        messages.error(request, 'Chỉ hỗ trợ file .xlsx')
        return redirect('draft_detail', draft_code=draft_code)

    try:
        wb = openpyxl.load_workbook(BytesIO(f.read()), read_only=True, data_only=True)
        ws = wb.active

        # Get next temp_id
        last = draft.persons.order_by('-temp_id').first()
        next_id = (last.temp_id + 1) if last else 1

        # Find header row
        headers = {}
        header_map = {
            'Họ tên': 'name', 'Ho ten': 'name', 'Tên': 'name', 'Ten': 'name', 'name': 'name',
            'Giới tính': 'gender', 'Gioi tinh': 'gender', 'GT': 'gender', 'gender': 'gender',
            'Năm sinh': 'birth_date', 'Nam sinh': 'birth_date', 'Ngày sinh': 'birth_date',
            'Ngay sinh': 'birth_date', 'birth_date': 'birth_date', 'Sinh': 'birth_date',
            'Năm mất': 'death_date', 'Nam mat': 'death_date', 'Ngày mất': 'death_date',
            'Ngay mat': 'death_date', 'death_date': 'death_date', 'Mất': 'death_date',
            'Ghi chú': 'notes', 'Ghi chu': 'notes', 'notes': 'notes',
            'Quan hệ': 'relation', 'Quan he': 'relation', 'relation': 'relation',
        }

        first_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
        for col_idx, cell_val in enumerate(first_row):
            if cell_val and str(cell_val).strip() in header_map:
                headers[header_map[str(cell_val).strip()]] = col_idx

        if 'name' not in headers:
            messages.error(request, 'Không tìm thấy cột "Họ tên" trong file Excel. Cần ít nhất 1 cột tên.')
            return redirect('draft_detail', draft_code=draft_code)

        gender_map = {
            'Nam': 'male', 'nam': 'male', 'M': 'male', 'm': 'male', 'male': 'male',
            'Nữ': 'female', 'nữ': 'female', 'Nu': 'female', 'nữ': 'female',
            'F': 'female', 'f': 'female', 'female': 'female',
        }

        created = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            name_val = row[headers['name']] if headers.get('name') is not None else None
            if not name_val or not str(name_val).strip():
                continue

            name = str(name_val).strip()
            gender_raw = str(row[headers['gender']]).strip() if 'gender' in headers and row[headers['gender']] else ''
            gender = gender_map.get(gender_raw, 'male')

            birth_date = str(row[headers['birth_date']]).strip() if 'birth_date' in headers and row[headers['birth_date']] else ''
            death_date = str(row[headers['death_date']]).strip() if 'death_date' in headers and row[headers['death_date']] else ''
            notes = str(row[headers['notes']]).strip() if 'notes' in headers and row[headers['notes']] else ''

            # Clean up "None" strings from empty cells
            if birth_date == 'None': birth_date = ''
            if death_date == 'None': death_date = ''
            if notes == 'None': notes = ''

            DraftPerson.objects.create(
                submission=draft,
                temp_id=next_id,
                name=name,
                gender=gender,
                birth_date=birth_date,
                death_date=death_date,
                notes=notes,
                relation_to_submitter='other',
            )
            next_id += 1
            created += 1

        wb.close()

        if created > 0:
            messages.success(request, f'Đã nhập {created} người từ Excel. Giờ hãy nối quan hệ!')
            return redirect('draft_relationship_builder', draft_code=draft_code)
        else:
            messages.warning(request, 'Không tìm thấy dữ liệu trong file Excel.')

    except Exception as e:
        messages.error(request, f'Lỗi đọc file Excel: {str(e)}')

    return redirect('draft_detail', draft_code=draft_code)


class DraftRelationshipBuilderView(TemplateView):
    """Trang kéo-nối quan hệ trực quan"""
    template_name = 'drafts/relationship_builder.html'

    def dispatch(self, request, *args, **kwargs):
        self.draft = get_object_or_404(DraftSubmission, draft_code=kwargs['draft_code'])
        if not _can_access_draft(request, self.draft):
            return HttpResponseForbidden('Không có quyền')
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['draft'] = self.draft
        persons = []
        for p in self.draft.persons.order_by('temp_id'):
            persons.append({
                'pk': p.pk,
                'temp_id': p.temp_id,
                'name': p.name,
                'gender': p.gender,
                'birth_date': p.birth_date,
                'death_date': p.death_date,
                'father_temp_id': p.father_temp_id,
                'mother_temp_id': p.mother_temp_id,
                'relation_to_submitter': p.relation_to_submitter,
                'photo_url': p.photo.url if p.photo else '',
                'notes': p.notes,
                'birth_order': p.birth_order,
                'chi_number': p.chi_number,
                'generation': p.generation,
            })
        ctx['persons'] = persons
        # Spouse relations as separate list
        spouse_rels = self.draft.spouse_relations.order_by('wife_order').values(
            'person1_temp_id', 'person2_temp_id', 'status', 'wife_order',
        )
        ctx['spouse_relations'] = list(spouse_rels)
        ctx['can_edit'] = _can_edit_draft(self.request, self.draft)
        return ctx


def draft_quick_add_person(request, draft_code):
    """AJAX: Thêm nhanh người từ relationship builder"""
    draft = get_object_or_404(DraftSubmission, draft_code=draft_code)
    if not _can_edit_draft(request, draft):
        return JsonResponse({'error': 'Không có quyền sửa'}, status=403)
    if request.method != 'POST':
        return JsonResponse({'error': 'POST only'}, status=405)

    name = request.POST.get('name', '').strip()
    gender = request.POST.get('gender', 'male')
    birth_date = request.POST.get('birth_date', '').strip()
    death_date = request.POST.get('death_date', '').strip()

    if not name:
        return JsonResponse({'error': 'Thiếu họ tên'}, status=400)
    if gender not in ('male', 'female'):
        return JsonResponse({'error': 'Giới tính không hợp lệ'}, status=400)
    if draft.persons.count() >= MAX_PERSONS_PER_DRAFT:
        return JsonResponse({'error': f'Tối đa {MAX_PERSONS_PER_DRAFT} người'}, status=400)

    last = draft.persons.order_by('-temp_id').first()
    temp_id = (last.temp_id + 1) if last else 1

    person = DraftPerson.objects.create(
        submission=draft,
        temp_id=temp_id,
        name=name,
        gender=gender,
        birth_date=birth_date,
        death_date=death_date,
        relation_to_submitter='other',
    )

    return JsonResponse({
        'ok': True,
        'person': {
            'pk': person.pk,
            'temp_id': person.temp_id,
            'name': person.name,
            'gender': person.gender,
            'birth_date': person.birth_date,
            'death_date': person.death_date,
            'father_temp_id': None,
            'mother_temp_id': None,
            'relation_to_submitter': 'other',
        }
    })


def draft_edit_person_ajax(request, draft_code, person_pk):
    """AJAX: Sửa thông tin người từ relationship builder"""
    draft = get_object_or_404(DraftSubmission, draft_code=draft_code)
    if not _can_edit_draft(request, draft):
        return JsonResponse({'error': 'Không có quyền sửa'}, status=403)
    person = get_object_or_404(DraftPerson, pk=person_pk, submission=draft)

    if request.method != 'POST':
        return JsonResponse({'error': 'POST only'}, status=405)

    person.name = request.POST.get('name', person.name).strip()
    person.gender = request.POST.get('gender', person.gender)
    person.birth_date = request.POST.get('birth_date', '').strip()
    person.death_date = request.POST.get('death_date', '').strip()
    person.notes = request.POST.get('notes', '').strip()
    person.relation_to_submitter = request.POST.get(
        'relation_to_submitter', person.relation_to_submitter
    )

    for fld in ('birth_order', 'chi_number', 'generation'):
        val = request.POST.get(fld, '').strip()
        setattr(person, fld, int(val) if val else None)

    # Handle father/mother changes
    for parent_field in ('father_temp_id', 'mother_temp_id'):
        val = request.POST.get(parent_field, '').strip()
        if val:
            parent_id = int(val)
            # Validate parent exists in same draft
            if draft.persons.filter(temp_id=parent_id).exists():
                setattr(person, parent_field, parent_id)
            # else: silently ignore invalid parent id
        else:
            setattr(person, parent_field, None)

    if 'photo' in request.FILES:
        person.photo = request.FILES['photo']

    if not person.name:
        return JsonResponse({'error': 'Thiếu họ tên'}, status=400)
    if person.gender not in ('male', 'female'):
        return JsonResponse({'error': 'Giới tính không hợp lệ'}, status=400)

    person.save()

    return JsonResponse({
        'ok': True,
        'person': {
            'pk': person.pk,
            'temp_id': person.temp_id,
            'name': person.name,
            'gender': person.gender,
            'birth_date': person.birth_date,
            'death_date': person.death_date,
            'father_temp_id': person.father_temp_id,
            'mother_temp_id': person.mother_temp_id,
            'relation_to_submitter': person.relation_to_submitter,
            'photo_url': person.photo.url if person.photo else '',
            'notes': person.notes,
            'birth_order': person.birth_order,
            'chi_number': person.chi_number,
            'generation': person.generation,
        }
    })


def draft_save_relation(request, draft_code):
    """AJAX: Lưu quan hệ giữa 2 người trong bản khai"""
    draft = get_object_or_404(DraftSubmission, draft_code=draft_code)
    if not _can_edit_draft(request, draft):
        return JsonResponse({'error': 'Không có quyền sửa'}, status=403)
    if request.method != 'POST':
        return JsonResponse({'error': 'POST only'}, status=405)

    from_id = request.POST.get('from_id')  # temp_id of source
    to_id = request.POST.get('to_id')      # temp_id of target
    rel_type = request.POST.get('rel_type')  # father, mother, spouse

    if not all([from_id, to_id, rel_type]):
        return JsonResponse({'error': 'Missing params'}, status=400)

    try:
        from_id = int(from_id)
        to_id = int(to_id)
    except ValueError:
        return JsonResponse({'error': 'Invalid IDs'}, status=400)

    person_from = draft.persons.filter(temp_id=from_id).first()
    person_to = draft.persons.filter(temp_id=to_id).first()

    if not person_from or not person_to:
        return JsonResponse({'error': 'Person not found'}, status=404)

    if rel_type == 'father':
        # "from" là cha của "to"
        person_to.father_temp_id = from_id
        person_to.save(update_fields=['father_temp_id', 'updated_at'])
    elif rel_type == 'mother':
        # "from" là mẹ của "to"
        person_to.mother_temp_id = from_id
        person_to.save(update_fields=['mother_temp_id', 'updated_at'])
    elif rel_type in ('spouse', 'spouse_divorced', 'spouse_widowed'):
        # Tạo/cập nhật DraftSpouseRelation
        status_map = {'spouse': 'married', 'spouse_divorced': 'divorced', 'spouse_widowed': 'widowed'}
        status = status_map[rel_type]
        lo, hi = min(from_id, to_id), max(from_id, to_id)
        # Auto wife_order: next number for this husband
        male_id = from_id if person_from.gender == 'male' else to_id
        existing_count = DraftSpouseRelation.objects.filter(
            submission=draft
        ).filter(
            models.Q(person1_temp_id=male_id) | models.Q(person2_temp_id=male_id)
        ).exclude(
            person1_temp_id=lo, person2_temp_id=hi
        ).count()
        rel, created = DraftSpouseRelation.objects.update_or_create(
            submission=draft, person1_temp_id=lo, person2_temp_id=hi,
            defaults={'status': status},
        )
        if created:
            rel.wife_order = existing_count + 1
            rel.save(update_fields=['wife_order'])
    elif rel_type == 'child':
        # "from" là con của "to" → to là cha/mẹ
        if person_to.gender == 'male':
            person_from.father_temp_id = to_id
        else:
            person_from.mother_temp_id = to_id
        person_from.save(update_fields=['father_temp_id', 'mother_temp_id', 'updated_at'])
    elif rel_type == 'remove':
        # Xóa quan hệ between from and to
        changed = False
        if person_to.father_temp_id == from_id:
            person_to.father_temp_id = None
            changed = True
        if person_to.mother_temp_id == from_id:
            person_to.mother_temp_id = None
            changed = True
        if person_from.father_temp_id == to_id:
            person_from.father_temp_id = None
            changed = True
        if person_from.mother_temp_id == to_id:
            person_from.mother_temp_id = None
            changed = True
        # Xóa spouse relation
        lo, hi = min(from_id, to_id), max(from_id, to_id)
        deleted = DraftSpouseRelation.objects.filter(
            submission=draft, person1_temp_id=lo, person2_temp_id=hi
        ).delete()[0]
        if deleted:
            changed = True
        if changed:
            person_from.save(update_fields=['father_temp_id', 'mother_temp_id', 'updated_at'])
            person_to.save(update_fields=['father_temp_id', 'mother_temp_id', 'updated_at'])
    else:
        return JsonResponse({'error': f'Unknown rel_type: {rel_type}'}, status=400)

    return JsonResponse({'ok': True})


# ═══════════════════════════════════════════════════════════════
# ADMIN VIEWS — Cần login + quyền editor
# ═══════════════════════════════════════════════════════════════

class EditorRequiredMixin(LoginRequiredMixin, UserPassesTestMixin):
    def test_func(self):
        return self.request.user.is_authenticated and self.request.user.can_edit_members()


class AdminDraftListView(EditorRequiredMixin, ListView):
    """Danh sách bản khai cho admin"""
    model = DraftSubmission
    template_name = 'drafts/admin_list.html'
    context_object_name = 'drafts'
    paginate_by = 20

    def get_queryset(self):
        qs = super().get_queryset().annotate(person_count=Count('persons'))
        status = self.request.GET.get('status')
        if status:
            qs = qs.filter(status=status)
        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['status_filter'] = self.request.GET.get('status', '')
        ctx['counts'] = {
            'submitted': DraftSubmission.objects.filter(status='submitted').count(),
            'draft': DraftSubmission.objects.filter(status='draft').count(),
            'approved': DraftSubmission.objects.filter(status='approved').count(),
            'rejected': DraftSubmission.objects.filter(status='rejected').count(),
        }
        return ctx


class AdminDraftReviewView(EditorRequiredMixin, DetailView):
    """Admin review chi tiết bản khai"""
    model = DraftSubmission
    template_name = 'drafts/admin_review.html'
    context_object_name = 'draft'
    slug_field = 'draft_code'
    slug_url_kwarg = 'draft_code'

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['persons'] = self.object.persons.order_by('temp_id')
        ctx['all_members'] = FamilyMember.objects.select_related('chi').order_by(
            'chi__number', 'generation', 'birth_order', 'name'
        )[:500]
        ctx['chi_list'] = Chi.objects.all()
        return ctx


def admin_draft_set_action(request, draft_code, person_pk):
    """Admin: set action (create_new / link_existing / skip) cho từng person"""
    if not (request.user.is_authenticated and request.user.can_edit_members()):
        return JsonResponse({'error': 'Unauthorized'}, status=403)
    if request.method != 'POST':
        return JsonResponse({'error': 'POST only'}, status=405)

    draft = get_object_or_404(DraftSubmission, draft_code=draft_code)
    person = get_object_or_404(DraftPerson, pk=person_pk, submission=draft)

    action = request.POST.get('action')
    linked_member_id = request.POST.get('linked_member')

    if action not in ('create_new', 'link_existing', 'skip', 'pending'):
        return JsonResponse({'error': 'Invalid action'}, status=400)

    person.action = action
    if action == 'link_existing' and linked_member_id:
        try:
            person.linked_member = FamilyMember.objects.get(pk=int(linked_member_id))
        except (FamilyMember.DoesNotExist, ValueError):
            return JsonResponse({'error': 'Member not found'}, status=404)
    else:
        person.linked_member = None

    person.save(update_fields=['action', 'linked_member', 'updated_at'])
    return JsonResponse({'ok': True, 'action': action, 'person_id': person.pk})


def admin_draft_approve(request, draft_code):
    """Admin: Duyệt & nhập bản khai vào FamilyMember chính thức"""
    if not (request.user.is_authenticated and request.user.can_edit_members()):
        return HttpResponseForbidden('Unauthorized')
    if request.method != 'POST':
        return redirect('admin_draft_review', draft_code=draft_code)

    draft = get_object_or_404(DraftSubmission, draft_code=draft_code)
    if draft.status == 'approved':
        messages.warning(request, 'Bản khai này đã được duyệt trước đó.')
        return redirect('admin_draft_review', draft_code=draft_code)

    persons = draft.persons.order_by('temp_id')

    # Check: tất cả person phải có action != pending
    pending = persons.filter(action='pending')
    if pending.exists():
        names = ', '.join(p.name for p in pending[:5])
        messages.error(request, f'Còn {pending.count()} người chưa xử lý: {names}. Hãy chọn hành động cho từng người.')
        return redirect('admin_draft_review', draft_code=draft_code)

    # Map temp_id → FamilyMember PK (for linking father/mother)
    temp_to_member = {}
    created_count = 0
    linked_count = 0
    skipped_count = 0

    # Pass 1: Process link_existing and skip
    for p in persons:
        if p.action == 'link_existing' and p.linked_member:
            temp_to_member[p.temp_id] = p.linked_member
            linked_count += 1
        elif p.action == 'skip':
            skipped_count += 1

    # Pass 2: Create new members (respecting dependency order — parents first)
    # Sort by generation/temp_id to create parents before children
    to_create = list(persons.filter(action='create_new').order_by('generation', 'temp_id'))

    for p in to_create:
        # Resolve chi
        chi_obj = None
        if p.chi_number:
            chi_obj = Chi.objects.filter(number=p.chi_number).first()

        # Resolve father/mother from temp_to_member
        father = temp_to_member.get(p.father_temp_id) if p.father_temp_id else None
        mother = temp_to_member.get(p.mother_temp_id) if p.mother_temp_id else None

        member = FamilyMember(
            name=p.name,
            gender=p.gender,
            chi=chi_obj,
            generation=p.generation,
            birth_date=p.birth_date or '',
            death_date=p.death_date or '',
            birth_order=p.birth_order,
            notes=p.notes,
            member_type='blood',
            created_by=request.user,
        )
        if father:
            member.father = father
            # Inherit chi and generation from father if not set
            if not member.chi and father.chi:
                member.chi = father.chi
            if not member.generation and father.generation is not None:
                member.generation = father.generation + 1
        if mother:
            member.mother = mother

        member.save()  # triggers auto person_code generation

        # Handle photo
        if p.photo:
            member.photo = p.photo
            member.save(update_fields=['photo'])

        temp_to_member[p.temp_id] = member
        p.linked_member = member
        p.save(update_fields=['linked_member', 'updated_at'])
        created_count += 1

    # Pass 3: Link spouses from DraftSpouseRelation
    for rel in draft.spouse_relations.all():
        m1 = temp_to_member.get(rel.person1_temp_id)
        m2 = temp_to_member.get(rel.person2_temp_id)
        if m1 and m2:
            # Determine husband/wife
            if m1.gender == 'male':
                husband, wife = m1, m2
            else:
                husband, wife = m2, m1
            SpouseRelation.objects.get_or_create(
                husband=husband, wife=wife,
                defaults={'status': rel.status, 'wife_order': rel.wife_order},
            )

    # Mark as approved
    draft.status = 'approved'
    draft.reviewed_by = request.user
    draft.reviewed_at = timezone.now()
    draft.review_notes = request.POST.get('review_notes', '')
    draft.save(update_fields=['status', 'reviewed_by', 'reviewed_at', 'review_notes', 'updated_at'])

    messages.success(
        request,
        f'Đã duyệt bản khai {draft.draft_code}: '
        f'tạo mới {created_count}, liên kết {linked_count}, bỏ qua {skipped_count} người.'
    )
    return redirect('admin_draft_review', draft_code=draft_code)


def admin_draft_reject(request, draft_code):
    """Admin: Từ chối bản khai"""
    if not (request.user.is_authenticated and request.user.can_edit_members()):
        return HttpResponseForbidden('Unauthorized')
    if request.method != 'POST':
        return redirect('admin_draft_review', draft_code=draft_code)

    draft = get_object_or_404(DraftSubmission, draft_code=draft_code)
    draft.status = 'rejected'
    draft.reviewed_by = request.user
    draft.reviewed_at = timezone.now()
    draft.review_notes = request.POST.get('review_notes', '')
    draft.save(update_fields=['status', 'reviewed_by', 'reviewed_at', 'review_notes', 'updated_at'])

    messages.success(request, f'Đã từ chối bản khai {draft.draft_code}.')
    return redirect('admin_draft_list')
