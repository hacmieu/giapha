"""
Convert Excel genealogy file to JSON for GoJS visualization
"""
import json
import re
from pathlib import Path
from openpyxl import load_workbook
from django.core.management.base import BaseCommand
from core.models import Genealogy, GenealogyPerson, GenealogySpouse, GenealogyRelationship


class Command(BaseCommand):
    help = 'Convert Excel genealogy file to JSON and save models'

    def add_arguments(self, parser):
        parser.add_argument('excel_file', type=str, help='Path to Excel file')
        parser.add_argument('--output', type=str, help='Output JSON file path')
        parser.add_argument('--genealogy-name', type=str, default='Gia Phả Họ Trần Chi 4',
                            help='Name of genealogy')

    def handle(self, *args, **options):
        excel_file = options['excel_file']
        output_file = options.get('output')
        genealogy_name = options['genealogy_name']

        import os
        if not os.path.exists(excel_file):
            self.stdout.write(self.style.ERROR(f'File not found: {excel_file}'))
            return

        self.stdout.write(f'Reading Excel file: {excel_file}')
        
        # Extract raw data
        raw_data = self.extract_excel_data(excel_file)
        self.stdout.write(self.style.SUCCESS(f'Extracted {len(raw_data)} records'))

        # Process data
        people_data, relationships = self.process_data(raw_data)
        self.stdout.write(self.style.SUCCESS(f'Processed {len(people_data)} unique people'))

        # Create genealogy object
        genealogy = self.save_to_database(genealogy_name, people_data, relationships)

        # Generate JSON
        json_data = self.generate_gojs_json(people_data, relationships)

        if output_file:
            with open(output_file, 'w', encoding='utf-8') as f:
                json.dump(json_data, f, ensure_ascii=False, indent=2)
            self.stdout.write(self.style.SUCCESS(f'Saved JSON to: {output_file}'))

        self.stdout.write(self.style.SUCCESS('Conversion completed successfully'))

    def extract_excel_data(self, excel_file):
        """Extract raw data from Excel"""
        wb = load_workbook(excel_file)
        ws = wb.active

        data = []
        current_generation = "Unknown"
        generation_counter = 0

        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=999, values_only=True), 1):
            first_col = row[0]

            if first_col is None:
                continue

            # Check if it's a generation header
            if isinstance(first_col, str) and 'Đời thứ' in first_col:
                current_generation = first_col.strip()
                # Extract generation number
                match = re.search(r'Đời thứ (\d+)', current_generation)
                if match:
                    generation_counter = int(match.group(1))
                continue

            # Skip header row
            if first_col in ['Mã & STT', 'Mã & STT ']:
                continue

            # Check if it's actual data (has ID code)
            if isinstance(first_col, str) and re.match(r'^\d{5}\.\d{4}$', str(first_col).strip()):
                record = {
                    'id': str(row[0]).strip(),
                    'name': str(row[1]).strip() if row[1] else '',
                    'birth_year': row[2],
                    'spouse_name': str(row[3]).strip() if row[3] else '',
                    'spouse_birth': row[4],
                    'father_id': str(row[5]).strip() if row[5] else None,
                    'location': str(row[6]).strip() if row[6] else '',
                    'notes': str(row[7]).strip() if row[7] else '',
                    'generation': current_generation,
                    'generation_num': generation_counter,
                    'row_num': row_idx
                }
                data.append(record)

        return data

    def process_data(self, raw_data):
        """Process raw data to consolidate multiple spouses"""
        people = {}
        relationships = []

        # First pass: create person entries
        for record in raw_data:
            person_id = record['id']

            if person_id not in people:
                # Extract display name (remove sequential number if exists)
                name = record['name']
                name_match = re.match(r'(\d{4})-(.+)', name)
                display_name = name_match.group(2) if name_match else name

                people[person_id] = {
                    'id': person_id,
                    'name': display_name,
                    'full_record_name': name,
                    'birth_year': record['birth_year'],
                    'location': record['location'],
                    'notes': record['notes'],
                    'generation': record['generation'],
                    'generation_num': record['generation_num'],
                    'spouses': [],
                    'father_id': record['father_id']
                }

            # Add spouse if exists
            if record['spouse_name']:
                spouse_data = self.parse_spouse_name(record['spouse_name'])
                people[person_id]['spouses'].extend(spouse_data)

        # Second pass: create relationships
        for person_id, person in people.items():
            if person['father_id']:
                # Try to resolve father_id to actual ID
                father_id = self.resolve_father_id(person['father_id'], people)
                if father_id:
                    relationships.append({
                        'from_id': father_id,
                        'to_id': person_id,
                        'type': 'parent'
                    })
                    person['father_id'] = father_id

        return people, relationships

    def parse_spouse_name(self, spouse_str):
        """Parse spouse name which might contain multiple spouses"""
        spouses = []
        # Handle numbered spouses like "1.KHÚC VĂN THƯ," or "2. VŨ THỊ LÊnh"
        lines = spouse_str.split(',')

        for line in lines:
            line = line.strip()
            if not line:
                continue

            # Remove numbering if exists
            match = re.match(r'^[\d]{1,2}\.\s*(.+)$', line)
            if match:
                name = match.group(1).strip()
            else:
                name = line

            if name and name.strip() != ".":
                spouses.append({'name': name.strip(), 'birth_year': None})

        return spouses

    def resolve_father_id(self, father_ref, people):
        """Try to resolve father reference to actual ID"""
        # 1. Already in exact ID format NNNNN.NNNN
        if re.match(r'^\d{5}\.\d{4}$', father_ref):
            return father_ref if father_ref in people else None

        # 2. Exact full_record_name or display name match
        for person_id, person in people.items():
            if person['full_record_name'] == father_ref or person['name'] == father_ref:
                return person_id

        # 3. Match by STT prefix alone (e.g. "0027" from "0027-Trần Đình Bàng")
        #    Handles typos, extra nicknames, case differences in column F
        stt_match = re.match(r'^(\d{4})-', father_ref)
        if stt_match:
            stt = stt_match.group(1)
            for person_id, person in people.items():
                if re.match(r'^' + stt + r'-', person.get('full_record_name', '')):
                    return person_id

        return None

    def save_to_database(self, genealogy_name, people_data, relationships):
        """Save to Django models"""
        # Create Genealogy object
        genealogy, created = Genealogy.objects.get_or_create(
            name=genealogy_name,
            defaults={
                'total_people': len(people_data),
                'description': 'Gia phả Họ Trần chi 4'
            }
        )

        if created:
            self.stdout.write(self.style.SUCCESS(f'Created Genealogy: {genealogy_name}'))
        else:
            # Clear existing data
            GenealogyPerson.objects.filter(genealogy=genealogy).delete()

        # Create Person entries
        for person_id, person in people_data.items():
            gp = GenealogyPerson.objects.create(
                genealogy=genealogy,
                person_id=person_id,
                name=person['name'],
                generation=person['generation_num'],
                birth_year=person['birth_year'],
                location=person['location'],
                notes=person['notes'],
                father_id=person['father_id']
            )

            # Create spouse entries
            for spouse in person['spouses']:
                GenealogySpouse.objects.create(
                    person=gp,
                    name=spouse['name'],
                    birth_year=spouse['birth_year']
                )

        # Create relationships
        GenealogyRelationship.objects.filter(genealogy=genealogy).delete()
        for rel in relationships:
            try:
                from_person = GenealogyPerson.objects.get(
                    genealogy=genealogy,
                    person_id=rel['from_id']
                )
                to_person = GenealogyPerson.objects.get(
                    genealogy=genealogy,
                    person_id=rel['to_id']
                )
                GenealogyRelationship.objects.create(
                    genealogy=genealogy,
                    from_person=from_person,
                    to_person=to_person,
                    relationship_type=rel['type']
                )
            except GenealogyPerson.DoesNotExist:
                pass

        return genealogy

    def generate_gojs_json(self, people_data, relationships):
        """Generate JSON in GoJS format"""
        # Create node data
        nodeDataArray = []
        for person_id, person in people_data.items():
            spouse_names = ', '.join([s['name'] for s in person['spouses']])

            node = {
                'key': person_id,
                'name': person['name'],
                'generation': person['generation_num'],
                'location': person['location'],
                'spouses': spouse_names,
                'notes': person['notes'],
                'birthYear': person['birth_year'],
                'fatherId': person['father_id']
            }
            nodeDataArray.append(node)

        # Create link data
        linkDataArray = []
        for rel in relationships:
            link = {
                'from': rel['from_id'],
                'to': rel['to_id'],
                'category': rel['type']
            }
            linkDataArray.append(link)

        return {
            'nodeDataArray': nodeDataArray,
            'linkDataArray': linkDataArray,
            'metadata': {
                'totalPeople': len(people_data),
                'totalRelationships': len(relationships),
                'generationRange': [4, 12]
            }
        }
